"""Generate the Dataset Cartography paper-and-reproduction presentation.

Run from the repository root:
    uv run --extra presentation python presentation/generate_presentation.py
"""
from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PRESENTATION = ROOT / "presentation"
ASSETS = PRESENTATION / "assets"
GENERATED = ASSETS / "generated"
RESULTS = ROOT / "results"
OUTPUT = PRESENTATION / "dataset_cartography_reproduction.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

BG = "F7F5F0"
TEXT = "202A35"
MUTED = "5F6B75"
PAPER = "697386"
OURS = "B65C3A"
TEAL = "2B6F6D"
GREEN = "4F7661"
AMBER = "B47A2A"
RED = "A94B46"
PALE = "ECE9E2"
WHITE = "FFFFFF"
GRID = "D8D4CC"
FONT = "Aptos"
MATH_FONT = "Cambria Math"

# These additions make the embedded notes a complete 20–25 minute research talk
# while keeping the visible slides concise. They add explanation, not new claims.
NOTE_EXTENSIONS = {
    "Title": "The structure is motivation, method, original evidence, reproduction protocol, results, and interpretation. Appendix slides hold configuration details so the main argument stays readable.",
    "Aggregate metrics hide example-level structure": "Two runs with identical final accuracy can contain very different datasets: one may have many consistently easy cases, while another may depend on unstable borderline cases. Aggregate evaluation therefore cannot tell us which examples drive learning, which are redundant, or which deserve annotation review.",
    "SNLI is a three-way natural-language inference task": "Entailment means the hypothesis must follow from the premise, contradiction means it cannot be true given the premise, and neutral covers cases where neither conclusion follows. The task is simple to state but exposes lexical, logical, and commonsense reasoning failures.",
    "Training dynamics record the path, not only the endpoint": "Recording once per epoch is computationally cheap because the model already performs these forward passes during training. The original work used all epochs, including the first, and reports that early approximations can still correlate well with converged maps.",
    "Three examples produce three distinct trajectories": "The unstable example also shows why a final prediction alone is insufficient: at the last epoch it looks confidently correct, but its earlier epochs tell us the model crossed a sharp decision boundary. Confidence and variability retain that history.",
    "Confidence summarizes how strongly an example is learned": "A high mean can coexist with one unusual epoch, and a low mean can arise from a bad annotation, a genuinely difficult case, or a model limitation. Confidence is therefore a diagnostic coordinate, not a verdict about data quality.",
    "Variability captures instability; correctness counts decisions": "The paper uses population standard deviation across epochs. Correctness takes only E plus one possible values, so it loses probability information, but it makes the distinction between stable correct and stable incorrect examples especially intuitive.",
    "A data map places every example by confidence and variability": "The authors observed broadly similar bell-shaped geometry across SNLI, MultiNLI, WinoGrande, and QNLI with RoBERTa-large. They also stress that changing the encoder can move individual coordinates even when the overall structure remains recognizable.",
    "Map regions become interpretable through real examples": "The middle example also warns against reading the gold label as unquestionable truth: practicing tennis is compatible with playing tennis, so contradiction is debatable. The map helps prioritize such cases for human inspection without asserting the correction automatically.",
    "The paper tested whether map regions predict useful behavior": "The subset experiments retrain from scratch after ranking examples, controlling the amount of training data. OOD sets were collected independently or designed adversarially, so they ask whether a selection learns something broader than the original dataset distribution.",
    "Ambiguous subsets often supported stronger generalization": "On WinoGrande, easy selections underperformed the random subset, while ambiguous and hard selections were competitive or better. For SNLI, the paper reports improvements across the Diagnostics categories, although the aggregate numbers shown here are the comparison we reproduce.",
    "Easy examples can be necessary for optimization": "This result suggests a curriculum-like balance: difficult variable examples may define useful distinctions, while some stable examples help the optimizer find a workable solution. The paper leaves the optimal mixture as an open question rather than prescribing one fixed ratio.",
    "Hard-to-learn regions are enriched for annotation problems": "The simple noise classifier used confidence as its single feature and achieved a sanity-check score of 100 percent F1 on similarly constructed artificial data. Real-data human review was weaker, which is expected because naturally difficult and mislabeled examples are not the same population.",
    "The selected SNLI experiment tests data efficiency and generalization": "Diagnostics contains categories for lexical semantics, predicate–argument structure, logic, and knowledge. We report one aggregate accuracy because that is what our evaluator produces. This keeps ID and OOD claims distinct and avoids implying that a higher SNLI score guarantees broader reasoning.",
    "Reproduction protocol follows the published SNLI selection": "Using released coordinates matters: it reproduces the paper's selection directly before asking whether our locally reconstructed map agrees. Otherwise a difference in final accuracy could combine selection error with training error and would be harder to interpret.",
    "Hardware constraints changed implementation details": "Gradient accumulation preserves optimizer batch frequency but not every floating-point operation, and BF16 changes numerical rounding. Dynamic padding preserves attention masks and the 128-token truncation limit while reducing wasted computation. These are practical, documented deviations.",
    "In-distribution accuracy closely matches the paper": "The ordering is not identical because our full run exceeds ambiguous by 0.34 points, whereas the published ambiguous value exceeds full by 0.2. Given best-of-three versus one seed, the defensible conclusion is near-equivalent performance with one third of the data, not that ambiguous selection always wins ID.",
    "The ambiguous subset also leads on NLI Diagnostics": "Here the qualitative ordering is especially informative: ambiguous is best in both paper and reproduction, and random is lowest. Hard-to-learn roughly matches the paper's OOD result even though its ID result is lower, reinforcing that ID and OOD answer different questions.",
    "Reproduced maps agree more on confidence than variability": "Set overlap is a stricter criterion than correlation because ranking changes near the one-third cutoff can swap membership. A 75.2 percent ambiguous overlap therefore coexists naturally with a Spearman variability correlation of 0.772. The excluded 0.09 percent is transparently reported.",
    "The selection result reproduces; map details remain sensitive": "Practical uses include dataset auditing, selecting compact training sets, and studying uncertainty, but computing the initial map still requires training on the full dataset. Our evidence is one dataset, one model family, and one seed, so broader deployment should validate the map again under the intended setup.",
}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, size: float, color: str = TEXT, bold: bool = False, font: str = FONT, italic: bool = False) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    run.font.italic = italic


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.notes: list[tuple[str, str]] = []

    def slide(self, title: str, *, section: str = "", citation: str = "", appendix: bool = False):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(BG)
        if section:
            textbox(slide, 0.55, 0.24, 3.0, 0.22, section.upper(), 8.5, OURS if section == "Our reproduction" else PAPER, bold=True, char_space=1.5)
        textbox(slide, 0.55, 0.52, 12.1, 0.48, title, 24, TEXT, bold=True)
        line(slide, 0.55, 1.12, 12.25, 1.12, GRID, 1)
        page = f"A{len(self.prs.slides) - 20}" if appendix else str(len(self.prs.slides))
        textbox(slide, 12.28, 7.14, 0.45, 0.16, page, 7.5, MUTED, align=PP_ALIGN.RIGHT)
        if citation:
            textbox(slide, 0.55, 7.10, 11.2, 0.22, citation, 7.2, MUTED)
        return slide

    def add_notes(self, slide, title: str, notes: str) -> None:
        clean = " ".join((notes + " " + NOTE_EXTENSIONS.get(title, "")).split())
        slide.notes_slide.notes_text_frame.text = clean
        self.notes.append((title, clean))

    def save(self) -> None:
        self.prs.save(OUTPUT)
        lines = ["# Speaker notes", "", "Target duration: 22–25 minutes for slides 1–20; appendix is backup.", ""]
        for i, (title, note) in enumerate(self.notes, 1):
            label = f"A{i - 20}" if i > 20 else str(i)
            lines.extend([f"## {label}. {title}", "", note, ""])
        (PRESENTATION / "speaker_notes.md").write_text("\n".join(lines), encoding="utf-8")


def textbox(slide, x, y, w, h, text, size=16, color=TEXT, *, bold=False, align=PP_ALIGN.LEFT,
            valign=MSO_ANCHOR.TOP, margin=0.04, font=FONT, italic=False, char_space=None,
            fill=None, line_color=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill or BG)
    shape.line.color.rgb = rgb(line_color or fill or BG)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    set_run(run, size, color, bold, font, italic)
    if char_space is not None:
        run.font._element.set("spc", str(int(char_space * 100)))
    return shape


def rich_box(slide, x, y, w, h, runs, *, fill=BG, line_color=BG, margin=0.08,
             valign=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color)
    tf = shape.text_frame; tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    for spec in runs:
        run = p.add_run(); run.text = spec[0]
        set_run(run, spec[1], spec[2] if len(spec) > 2 else TEXT, spec[3] if len(spec) > 3 else False,
                spec[4] if len(spec) > 4 else FONT, spec[5] if len(spec) > 5 else False)
    return shape


def bullet_list(slide, x, y, w, h, items, *, size=16, color=TEXT, bullet_color=OURS, gap=9, fill=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill); shape.line.color.rgb = rgb(fill)
    tf = shape.text_frame; tf.clear()
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = 0
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        p.text = ""
        r = p.add_run(); r.text = "●  "; set_run(r, size * 0.58, bullet_color, bold=True)
        r = p.add_run(); r.text = item; set_run(r, size, color)
    return shape


def line(slide, x1, y1, x2, y2, color=GRID, width=1.5, dash=None):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color); shape.line.width = Pt(width)
    if dash is not None:
        shape.line.dash_style = dash
    return shape


def pill(slide, x, y, w, text, color, *, fill=BG):
    return textbox(slide, x, y, w, 0.32, text.upper(), 8.5, color, bold=True, align=PP_ALIGN.CENTER,
                   valign=MSO_ANCHOR.MIDDLE, fill=fill, line_color=color, radius=True)


def panel(slide, x, y, w, h, *, fill=WHITE, border=GRID):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill); shape.line.color.rgb = rgb(border)
    shape.line.width = Pt(0.8)
    return shape


def add_image_fit(slide, path: Path, x, y, w, h, *, contain=True):
    with Image.open(path) as im:
        iw, ih = im.size
    image_ratio = iw / ih; box_ratio = w / h
    if contain:
        if image_ratio > box_ratio:
            new_w, new_h = w, w / image_ratio
        else:
            new_h, new_w = h, h * image_ratio
        return slide.shapes.add_picture(str(path), Inches(x + (w - new_w) / 2), Inches(y + (h - new_h) / 2),
                                        width=Inches(new_w), height=Inches(new_h))
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def metric_table(slide, x, y, w, h, headers, rows, widths=None, *, paper_cols=(), ours_cols=(), font_size=12):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    for c, header in enumerate(headers):
        cell = table.cell(0, c); cell.text = header
        cell.fill.solid(); cell.fill.fore_color.rgb = rgb(PALE)
        cell.margin_left = cell.margin_right = Inches(0.06)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            for run in p.runs: set_run(run, font_size - 1, TEXT, bold=True)
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            cell = table.cell(r, c); cell.text = str(value)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(WHITE if r % 2 else "F2F0EA")
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                for run in p.runs:
                    color = PAPER if c in paper_cols else OURS if c in ours_cols else TEXT
                    set_run(run, font_size, color, bold=c in ours_cols)
    return table_shape


def save_plot(path: Path) -> None:
    plt.savefig(path, dpi=190, bbox_inches="tight", facecolor=f"#{BG}")
    plt.close()


def make_assets(metrics: pd.DataFrame, trajectories: pd.DataFrame) -> None:
    GENERATED.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 12, "axes.titlesize": 15,
        "axes.labelsize": 12, "axes.edgecolor": "#5F6B75", "axes.facecolor": f"#{BG}",
        "figure.facecolor": f"#{BG}", "grid.color": "#D8D4CC", "grid.alpha": 0.6,
    })

    # One illustrative trajectory for the conceptual shift slide.
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    epochs = np.arange(1, 7); probs = [.42, .58, .76, .81, .88, .91]
    ax.plot(epochs, probs, color="#2B6F6D", linewidth=3, marker="o", markersize=7)
    for e, p in zip(epochs, probs): ax.text(e, p + .055, f"{p:.2f}", ha="center", fontsize=10)
    ax.set(xlabel="Epoch", ylabel="P(gold label)", ylim=(0, 1.05), xticks=epochs)
    ax.grid(axis="y"); ax.spines[["top", "right"]].set_visible(False)
    save_plot(GENERATED / "single_trajectory.png")

    # Actual trajectories recorded by this reproduction.
    fig, ax = plt.subplots(figsize=(8.4, 4.1))
    labels = {"easy": "stable high", "ambiguous": "unstable", "hard": "stable low"}
    colors = {"easy": "#4F7661", "ambiguous": "#B47A2A", "hard": "#A94B46"}
    for region, group in trajectories.groupby("region", sort=False):
        ax.plot(group.epoch, group.gold_probability, marker="o", linewidth=3, markersize=6,
                label=labels[region], color=colors[region])
    ax.set(xlabel="Epoch", ylabel="P(gold label)", ylim=(-.03, 1.07), xticks=range(1, 6))
    ax.grid(axis="y"); ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, ncols=3, loc="lower center", bbox_to_anchor=(.5, 1.01))
    save_plot(GENERATED / "three_trajectories.png")

    # Confidence and variability formula images.
    formulas = {
        "confidence_formula.png": r"$\mu_i = \frac{1}{E}\sum_{e=1}^{E}p_{\theta_e}(y_i\mid x_i)$",
        "variability_formula.png": r"$\sigma_i = \sqrt{\frac{1}{E}\sum_{e=1}^{E}\left(p_{\theta_e}(y_i\mid x_i)-\mu_i\right)^2}$",
    }
    for name, formula in formulas.items():
        fig = plt.figure(figsize=(8.0, 1.3), facecolor=f"#{BG}")
        fig.text(.5, .5, formula, ha="center", va="center", fontsize=28, color="#202A35")
        plt.axis("off"); save_plot(GENERATED / name)

    # Honest paired-dot comparisons with explicitly marked zoomed axes.
    for kind, pcol, ocol, xlim in (
        ("id", "paper_snli", "ours_snli", (89.5, 93.0)),
        ("ood", "paper_ood", "ours_ood", (59.5, 65.0)),
    ):
        fig, ax = plt.subplots(figsize=(8.9, 4.8))
        y = np.arange(len(metrics))[::-1]
        for yi, (_, row) in zip(y, metrics.iterrows()):
            ax.plot([row[pcol], row[ocol]], [yi, yi], color="#B8B4AC", linewidth=2, zorder=1)
            ax.scatter(row[pcol], yi, s=90, color="#697386", label="Paper" if yi == y[0] else None, zorder=2)
            ax.scatter(row[ocol], yi, s=90, color="#B65C3A", label="Our reproduction" if yi == y[0] else None, zorder=2)
            ax.text(row[pcol] - .05, yi + .18, f"{row[pcol]:.1f}", ha="right", color="#697386", fontsize=10)
            ax.text(row[ocol] + .05, yi - .22, f"{row[ocol]:.2f}", ha="left", color="#B65C3A", fontsize=10, weight="bold")
        ax.set_yticks(y, metrics.label)
        ax.set_xlim(*xlim); ax.set_xlabel("Accuracy (%) — zoomed axis")
        ax.grid(axis="x"); ax.spines[["top", "right", "left"]].set_visible(False)
        ax.legend(frameon=False, ncols=2, loc="lower center", bbox_to_anchor=(.5, 1.01))
        save_plot(GENERATED / f"{kind}_paired_dots.png")


def build_deck() -> None:
    metrics = pd.read_csv(RESULTS / "reproduction_metrics.csv")
    trajectories = pd.read_csv(ASSETS / "example_trajectories.csv")
    map_stats = json.loads((RESULTS / "map_comparison.json").read_text())
    make_assets(metrics, trajectories)
    deck = Deck()

    # 1 — Title
    slide = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = rgb(BG)
    line(slide, .75, 1.1, 2.4, 1.1, OURS, 4)
    textbox(slide, .75, 1.35, 11.7, 1.25, "Dataset Cartography: Mapping and\nDiagnosing Datasets with Training Dynamics", 30, TEXT, bold=True)
    textbox(slide, .78, 2.88, 10.8, .48, "Reproduction of Swayamdipta et al., EMNLP 2020", 17, PAPER)
    textbox(slide, .78, 4.52, 4.8, .38, "Julia Mazine", 16, TEXT, bold=True)
    textbox(slide, .78, 5.0, 8.0, .3, "RoBERTa-large · SNLI · NLI Diagnostics", 12, MUTED)
    textbox(slide, .78, 6.75, 8.8, .25, "github.com/JuliaMazine/dataset-cartography-reproduction", 9, MUTED)
    textbox(slide, 11.75, 6.75, .7, .25, "1 / 20", 8, MUTED, align=PP_ALIGN.RIGHT)
    deck.add_notes(slide, "Title", "This talk explains the Dataset Cartography paper first, then a focused reproduction of its SNLI data-selection result. The question running through the talk is whether the path a model takes while learning an example can reveal useful structure in the dataset. I will spend roughly the first half on the paper and the second half on what we reproduced.")

    # 2 — Motivation
    title = "Aggregate metrics hide example-level structure"
    slide = deck.slide(title, section="Paper")
    textbox(slide, .72, 1.42, 4.25, .65, "10,000 training examples", 22, TEXT, bold=True)
    # dot field
    for row in range(10):
        for col in range(18):
            color = PALE
            if (row * 18 + col) in {4, 17, 39, 58, 82, 95, 111, 137, 155, 169}: color = OURS
            elif (row * 18 + col) in {8, 28, 50, 73, 104, 126, 149}: color = TEAL
            shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(.78 + col * .22), Inches(2.2 + row * .22), Inches(.08), Inches(.08))
            shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color); shape.line.color.rgb = rgb(color)
    textbox(slide, 5.2, 1.55, .9, .4, "→", 28, MUTED, bold=True, align=PP_ALIGN.CENTER)
    panel(slide, 6.25, 1.48, 2.25, 1.35, fill=WHITE)
    textbox(slide, 6.42, 1.7, 1.9, .32, "Accuracy", 13, MUTED, align=PP_ALIGN.CENTER)
    textbox(slide, 6.42, 2.05, 1.9, .5, "92.0%", 27, PAPER, bold=True, align=PP_ALIGN.CENTER)
    panel(slide, 8.8, 1.48, 2.25, 1.35, fill=WHITE)
    textbox(slide, 8.97, 1.7, 1.9, .32, "Loss", 13, MUTED, align=PP_ALIGN.CENTER)
    textbox(slide, 8.97, 2.05, 1.9, .5, "0.21", 27, PAPER, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 6.25, 3.45, 5.8, .62, "These summaries say little about which examples were learned immediately, learned unreliably, or never learned.", 18, TEXT, bold=True)
    textbox(slide, 6.25, 4.65, 5.6, .85, "Can we diagnose a dataset by watching how the model learns each example?", 22, OURS, bold=True)
    deck.add_notes(slide, title, "We normally describe a training run with aggregate loss and evaluation accuracy. Those are necessary, but they collapse thousands of different example histories into one number. Some examples may become correct almost immediately, some may fluctuate, and some may resist the provided label. Dataset Cartography starts from the idea that this variation is useful information about the dataset itself.")

    # 3 — SNLI
    title = "SNLI is a three-way natural-language inference task"
    slide = deck.slide(title, section="Paper", citation="SNLI: Bowman et al. (2015). Examples shown are real SNLI training pairs.")
    panel(slide, .72, 1.38, 11.9, .82, fill="F0EEE8")
    textbox(slide, .92, 1.55, 1.1, .24, "PREMISE", 8.5, MUTED, bold=True)
    textbox(slide, 2.02, 1.47, 9.9, .45, '“Cafe Express” sign covered in graffiti.', 19, TEXT, bold=True)
    labels = [
        ("ENTAILMENT", "A sign covered in graffiti.", GREEN, "must be true"),
        ("NEUTRAL", "The sign is hard to read.", AMBER, "may be true"),
        ("CONTRADICTION", "The sign is brand new and clean.", RED, "cannot be true"),
    ]
    for i, (lab, hyp, color, desc) in enumerate(labels):
        y = 2.62 + i * 1.18
        panel(slide, .92, y, 10.9, .9, fill=WHITE, border=color)
        pill(slide, 1.12, y + .2, 1.65, lab, color, fill=WHITE)
        textbox(slide, 3.05, y + .14, 6.6, .34, hyp, 16, TEXT, bold=True)
        textbox(slide, 9.72, y + .18, 1.75, .26, desc, 11, color, align=PP_ALIGN.RIGHT)
    textbox(slide, 1.0, 6.32, 10.8, .42, "The model receives a premise–hypothesis pair and predicts their logical relation.", 17, TEXT, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "SNLI is central to our reproduction. Each example contains a premise and a hypothesis, and the task is to decide whether the hypothesis follows from the premise, could be true without following, or conflicts with it. These three hypotheses share one real SNLI premise. In-distribution evaluation uses held-out SNLI examples; later we will distinguish that from the hand-built NLI Diagnostics set.")

    # 4 — key shift
    title = "Training dynamics record the path, not only the endpoint"
    slide = deck.slide(title, section="Paper", citation="Swayamdipta et al. (2020), §2.1.")
    panel(slide, .74, 1.4, 4.25, 4.95, fill=WHITE)
    textbox(slide, 1.0, 1.7, 3.7, .3, "Ordinary analysis", 13, PAPER, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.02, 2.3, 3.65, .65, "What did the final model predict?", 23, TEXT, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.05, 3.42, 3.6, .42, "Epoch 6", 12, MUTED, align=PP_ALIGN.CENTER)
    panel(slide, 1.55, 4.0, 2.6, 1.0, fill="E7EFEA", border=GREEN)
    textbox(slide, 1.65, 4.17, 2.4, .24, "GOLD LABEL", 8.5, GREEN, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.65, 4.48, 2.4, .34, "0.91", 25, GREEN, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.1, 5.55, 3.55, .38, "one decision", 15, PAPER, italic=True, align=PP_ALIGN.CENTER)
    textbox(slide, 5.28, 3.32, .65, .5, "→", 30, OURS, bold=True, align=PP_ALIGN.CENTER)
    panel(slide, 6.1, 1.4, 6.25, 4.95, fill=WHITE)
    textbox(slide, 6.42, 1.7, 5.6, .3, "Dataset Cartography", 13, OURS, bold=True, align=PP_ALIGN.CENTER)
    add_image_fit(slide, GENERATED / "single_trajectory.png", 6.55, 2.02, 5.35, 3.15)
    textbox(slide, 6.52, 5.55, 5.45, .38, "a trajectory across training", 15, OURS, italic=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "The conceptual shift is simple. Instead of keeping only the final prediction, record the probability assigned to the gold label at the end of every epoch. The sequence of probabilities is an example’s training trajectory. The illustrated example rises smoothly from 0.42 to 0.91. Dataset Cartography summarizes each such trajectory and plots every training example.")

    # 5 — trajectories
    title = "Three examples produce three distinct trajectories"
    slide = deck.slide(title, section="Our reproduction", citation="Actual epoch-level probabilities from our full SNLI run; five logged epochs.")
    add_image_fit(slide, GENERATED / "three_trajectories.png", .72, 1.36, 7.35, 4.7)
    examples = [
        ("Stable high", "contradiction", "Beach children  →  men in a jungle", GREEN),
        ("Unstable", "contradiction", "Woman playing tennis  →  woman practices tennis", AMBER),
        ("Stable low", "entailment", "Girl photographed  →  boy photographed", RED),
    ]
    for i, (name, label, short, color) in enumerate(examples):
        y = 1.55 + i * 1.55
        textbox(slide, 8.28, y, 3.95, .3, name, 15, color, bold=True)
        textbox(slide, 8.28, y + .37, 3.95, .48, short, 12.5, TEXT)
        pill(slide, 8.28, y + .95, 1.25, label, color, fill=BG)
    textbox(slide, 1.2, 6.28, 10.8, .35, "How can these trajectories be summarized with a few comparable numbers?", 19, OURS, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "These are not stylized curves; they come from our recorded SNLI training dynamics. One example is assigned essentially full gold probability in every epoch. The tennis pair switches abruptly from near zero to near one, showing high instability. The girl-versus-boy example stays near zero because its provided entailment label conflicts with the sentences. These trajectories motivate confidence, variability, and correctness.")

    # 6 — confidence
    title = "Confidence summarizes how strongly an example is learned"
    slide = deck.slide(title, section="Paper", citation="Swayamdipta et al. (2020), Eq. 1 and §2.1.")
    add_image_fit(slide, GENERATED / "confidence_formula.png", 1.0, 1.48, 6.7, 1.35)
    textbox(slide, 1.12, 2.78, 6.5, .55, "Average probability assigned to the gold label throughout training.", 20, TEXT, bold=True, align=PP_ALIGN.CENTER)
    panel(slide, 8.05, 1.48, 4.25, 2.05, fill=WHITE)
    textbox(slide, 8.36, 1.72, 3.7, .26, "WHERE", 9, MUTED, bold=True)
    textbox(slide, 8.36, 2.12, 3.7, .3, "E   number of epochs", 14, TEXT)
    textbox(slide, 8.36, 2.55, 3.7, .3, "yᵢ   gold label", 14, TEXT)
    textbox(slide, 8.36, 2.98, 3.7, .3, "pθₑ  model probability at epoch e", 14, TEXT)
    line(slide, 1.0, 4.04, 12.2, 4.04, GRID, 1)
    panel(slide, 1.0, 4.42, 5.25, 1.45, fill="E7EFEA", border=GREEN)
    textbox(slide, 1.25, 4.67, 4.75, .35, "HIGH CONFIDENCE", 11, GREEN, bold=True)
    textbox(slide, 1.25, 5.12, 4.7, .48, "The model generally supports the provided label.", 17, TEXT)
    panel(slide, 7.0, 4.42, 5.25, 1.45, fill="F3E8E5", border=RED)
    textbox(slide, 7.25, 4.67, 4.75, .35, "LOW CONFIDENCE", 11, RED, bold=True)
    textbox(slide, 7.25, 5.12, 4.7, .48, "The model repeatedly struggles with that label.", 17, TEXT)
    textbox(slide, 1.05, 6.42, 11.25, .3, "Confidence is relative to the gold label and to this model and training procedure.", 13, MUTED, italic=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "Confidence is the mean gold-label probability over epochs. High confidence means the learner usually assigns strong probability to the annotation; low confidence means it repeatedly does not. This is not the model’s maximum class probability and it is not an objective property of the sentence pair. It depends on the gold label, the model family, and training procedure.")

    # 7 — variability and correctness
    title = "Variability captures instability; correctness counts decisions"
    slide = deck.slide(title, section="Paper", citation="Swayamdipta et al. (2020), §2.1.")
    add_image_fit(slide, GENERATED / "variability_formula.png", .85, 1.42, 7.15, 1.35)
    textbox(slide, 1.1, 2.72, 6.65, .5, "Standard deviation of gold-label confidence across epochs.", 19, TEXT, bold=True, align=PP_ALIGN.CENTER)
    panel(slide, 8.25, 1.47, 3.95, 1.82, fill=WHITE)
    textbox(slide, 8.52, 1.72, 3.4, .3, "CORRECTNESS", 11, TEAL, bold=True)
    textbox(slide, 8.52, 2.13, 3.35, .7, "Fraction of epochs in which the predicted class equals the gold label.", 16, TEXT)
    states = [
        ("stable high confidence", "low variability", GREEN, [0.85, .91, .90, .94, .93]),
        ("unstable confidence", "high variability", AMBER, [.12, .85, .20, .91, .38]),
        ("stable low confidence", "low variability", RED, [.06, .03, .05, .04, .06]),
    ]
    for i, (a, b, color, vals) in enumerate(states):
        y = 3.7 + i * .93
        textbox(slide, .92, y, 2.8, .28, a, 13, TEXT, bold=True)
        # sparkline in native vector
        for j in range(4):
            x1 = 4.0 + j * .46; x2 = 4.0 + (j + 1) * .46
            yy1 = y + .56 - vals[j] * .5; yy2 = y + .56 - vals[j + 1] * .5
            line(slide, x1, yy1, x2, yy2, color, 2.5)
        textbox(slide, 6.25, y, 2.4, .28, b, 13, color, bold=True)
    textbox(slide, 9.1, 4.22, 2.5, .92, "Variability can be low when the model is consistently right or consistently wrong.", 16, TEXT, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 9.1, 5.42, 2.5, .72, "Correctness helps distinguish those cases.", 15, TEAL, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "Variability is the standard deviation of those probabilities. An example can have low variability for two opposite reasons: it is confidently correct every epoch, or confidently inconsistent with its annotation every epoch. Correctness provides a coarser count of how often the predicted class matches the label. Together the measures distinguish stable-high, unstable, and stable-low behavior.")

    # 8 — map
    title = "A data map places every example by confidence and variability"
    slide = deck.slide(title, section="Paper", citation="Method: Swayamdipta et al. (2020), §2.2. Plot shown: our reproduced SNLI map.")
    add_image_fit(slide, RESULTS / "figures/snli_data_map.png", .6, 1.28, 8.0, 5.65)
    # region labels on right, avoiding arbitrary partitions
    pill(slide, 9.0, 1.63, 2.3, "easy-to-learn", GREEN, fill=BG)
    textbox(slide, 8.95, 2.05, 3.0, .58, "high confidence\nlow variability", 15, TEXT, align=PP_ALIGN.CENTER)
    pill(slide, 9.0, 3.15, 2.3, "ambiguous", AMBER, fill=BG)
    textbox(slide, 8.95, 3.57, 3.0, .58, "high variability\nconfidence may vary", 15, TEXT, align=PP_ALIGN.CENTER)
    pill(slide, 9.0, 4.67, 2.3, "hard-to-learn", RED, fill=BG)
    textbox(slide, 8.95, 5.09, 3.0, .58, "low confidence\noften low correctness", 15, TEXT, align=PP_ALIGN.CENTER)
    textbox(slide, 8.8, 6.15, 3.4, .5, "Every point is one training example.", 16, OURS, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "The data map uses variability on the horizontal axis and confidence on the vertical axis. Each point is one training example, and color shows the number of epochs predicted correctly. The bell shape gives three characteristic regions. The labels are descriptive and model-relative; they are not fixed semantic categories or permanent properties of an example.")

    # 9 — examples
    title = "Map regions become interpretable through real examples"
    slide = deck.slide(title, section="Our reproduction", citation="Examples and coordinates from our reproduced SNLI map; five logged epochs.")
    cards = [
        ("EASY-TO-LEARN", GREEN, "Two children are on the beach and play in wet sand.", "Two men are walking through a jungle.", "contradiction", "confidence 1.000  ·  variability < .001"),
        ("AMBIGUOUS", AMBER, "A woman in red and white clothing is playing tennis.", "Dressed all in white, the woman practices tennis.", "contradiction", "confidence .406  ·  variability .482"),
        ("HARD-TO-LEARN", RED, "Photographers photograph a girl sitting in a street.", "The photographer is taking a picture of a boy.", "entailment", "confidence < .001  ·  variability < .001"),
    ]
    for i, (region, color, premise, hyp, label, stat) in enumerate(cards):
        x = .64 + i * 4.2
        panel(slide, x, 1.42, 3.85, 4.95, fill=WHITE, border=color)
        pill(slide, x + .24, 1.7, 1.78, region, color, fill=WHITE)
        textbox(slide, x + .24, 2.26, 3.35, .25, "PREMISE", 8.5, MUTED, bold=True)
        textbox(slide, x + .24, 2.56, 3.34, .84, premise, 14, TEXT, bold=True)
        textbox(slide, x + .24, 3.58, 3.35, .25, "HYPOTHESIS", 8.5, MUTED, bold=True)
        textbox(slide, x + .24, 3.88, 3.34, .82, hyp, 14, TEXT, bold=True)
        pill(slide, x + .24, 4.94, 1.25, label, color, fill=WHITE)
        textbox(slide, x + .24, 5.54, 3.34, .48, stat, 11.5, color, bold=True)
    textbox(slide, .72, 6.58, 11.85, .3, "Ambiguous ≠ mislabeled. Hard-to-learn ≠ automatically wrong.", 16, TEXT, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "These examples make the geometry concrete. The beach-versus-jungle contradiction is consistently easy. The tennis pair is arguably compatible rather than contradictory, and the model switches sharply during training. The hard example’s entailment label says a girl entails a boy, so its near-zero confidence is unsurprising. These examples illustrate tendencies, not equivalences: ambiguity does not mean annotation error, and hard-to-learn does not guarantee one.")

    # 10 — evaluation program
    title = "The paper tested whether map regions predict useful behavior"
    slide = deck.slide(title, section="Paper", citation="Swayamdipta et al. (2020), §§3–6.")
    datasets = [("SNLI", "natural-language inference"), ("MultiNLI", "multi-genre inference"), ("WinoGrande", "commonsense reasoning"), ("QNLI", "question answering reformulated as NLI")]
    for i, (name, desc) in enumerate(datasets):
        x = .75 + i * 3.02
        textbox(slide, x, 1.48, 2.65, .42, name, 18, PAPER, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, x, 1.95, 2.65, .65, desc, 12.5, MUTED, align=PP_ALIGN.CENTER)
    line(slide, 1.0, 2.92, 12.0, 2.92, GRID, 1)
    questions = [
        ("DATA USEFULNESS", "Which regions support ID and OOD generalization?", TEAL),
        ("OPTIMIZATION", "Can a model learn from only ambiguous examples?", AMBER),
        ("ANNOTATION QUALITY", "Are low-confidence regions enriched for problems?", RED),
        ("UNCERTAINTY", "How do dynamics relate to human agreement?", PAPER),
    ]
    for i, (head, body, color) in enumerate(questions):
        y = 3.3 + i * .82
        textbox(slide, .9, y, 2.4, .28, head, 10, color, bold=True)
        textbox(slide, 3.25, y - .04, 8.6, .38, body, 17, TEXT)
    textbox(slide, 1.0, 6.65, 11.25, .28, "The maps were hypotheses about dataset structure; these experiments tested their meaning.", 16, OURS, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "The paper did more than visualize four datasets. It treated the regions as empirical hypotheses. The authors trained new models on selected regions, tested in-distribution and independently collected out-of-distribution sets, varied subset size to study optimization, injected label noise, and compared dynamics with human agreement. This validation program is why the method is interesting.")

    # 11 — ambiguous generalization
    title = "Ambiguous subsets often supported stronger generalization"
    slide = deck.slide(title, section="Paper", citation="Published values: Swayamdipta et al. (2020), Tables 2 and 3; best of 3 seeds for SNLI.")
    textbox(slide, .72, 1.37, 5.7, .45, "WinoGrande — 33% subsets", 18, TEXT, bold=True)
    metric_table(slide, .72, 1.88, 5.7, 3.05,
                 ["Training data", "ID", "OOD"],
                 [("Full 100%", "79.7", "86.0"), ("Random", "73.3", "85.6"), ("Hard-to-learn", "77.9", "87.2"), ("Ambiguous", "78.7", "87.6")],
                 widths=[2.8, 1.4, 1.4], paper_cols=(1, 2), font_size=13)
    textbox(slide, 6.88, 1.37, 5.7, .45, "SNLI — 33% ambiguous", 18, TEXT, bold=True)
    panel(slide, 6.88, 1.88, 5.7, 3.05, fill=WHITE)
    textbox(slide, 7.25, 2.2, 4.95, .3, "IN-DISTRIBUTION", 10, PAPER, bold=True, align=PP_ALIGN.CENTER)
    rich_box(slide, 7.25, 2.62, 4.95, .55, [("92.2%", 28, PAPER, True)], fill=WHITE, align=PP_ALIGN.CENTER)
    textbox(slide, 7.25, 3.26, 4.95, .3, "within 0.2 pp of full data", 13, MUTED, align=PP_ALIGN.CENTER)
    textbox(slide, 7.25, 3.82, 4.95, .3, "NLI DIAGNOSTICS", 10, PAPER, bold=True, align=PP_ALIGN.CENTER)
    rich_box(slide, 7.25, 4.17, 4.95, .55, [("63.5%", 28, PAPER, True)], fill=WHITE, align=PP_ALIGN.CENTER)
    textbox(slide, .82, 5.38, 11.65, .62, "Across these experiments, challenging and variable examples could preserve ID accuracy while improving independently collected OOD performance.", 20, TEXT, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.2, 6.28, 10.8, .36, "A recurring empirical pattern — not a universal law about every dataset or model.", 14, OURS, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "On WinoGrande, the most ambiguous third produced the strongest reported OOD score, 87.6, while retaining much more ID performance than a random third. In SNLI, the ambiguous third was within 0.2 percentage points of full-data ID performance and exceeded full-data Diagnostics accuracy. The paper found related patterns across datasets, but we should state this as empirical evidence for these settings, not a universal rule.")

    # 12 — easy optimization
    title = "Easy examples can be necessary for optimization"
    slide = deck.slide(title, section="Paper", citation="Swayamdipta et al. (2020), §4 and Figure 3.")
    # Left causal sequence
    panel(slide, .78, 1.42, 5.55, 4.95, fill=WHITE)
    textbox(slide, 1.08, 1.72, 4.95, .4, "WinoGrande: only the most ambiguous data", 17, TEXT, bold=True, align=PP_ALIGN.CENTER)
    for i, (pct, status, color) in enumerate([("25–50%", "learns", GREEN), ("≤17%", "chance-level", RED)]):
        y = 2.45 + i * 1.12
        textbox(slide, 1.25, y, 1.45, .42, pct, 22, color, bold=True, align=PP_ALIGN.CENTER)
        line(slide, 2.85, y + .2, 3.65, y + .2, GRID, 2)
        textbox(slide, 3.85, y, 1.7, .42, status, 18, color, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.15, 4.84, 4.8, .88, "With 17% ambiguous data, replacing only one tenth with easy examples restored learning and beat the random baseline on ID accuracy.", 16, TEXT, align=PP_ALIGN.CENTER)
    # Right nuance
    textbox(slide, 6.9, 1.62, 5.2, .45, "Different regions can serve different roles", 20, OURS, bold=True)
    roles = [
        ("ambiguous", "informative decision boundaries", AMBER),
        ("easy", "stable optimization signal", GREEN),
        ("too much easy", "weaker OOD performance", PAPER),
    ]
    for i, (name, role, color) in enumerate(roles):
        y = 2.48 + i * 1.08
        pill(slide, 7.0, y, 1.45, name, color, fill=BG)
        textbox(slide, 8.8, y + .02, 3.1, .36, role, 16, TEXT)
    textbox(slide, 6.95, 5.74, 5.05, .5, "“Ambiguous good, easy bad” would misstate the paper.", 18, RED, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "The optimization experiment adds an important qualification. WinoGrande models trained on 25 percent or more of the most ambiguous data learned successfully, but at 17 percent or less they collapsed to chance despite restarts. Replacing just one tenth of that 17 percent subset with easy examples restored learning and beat the random ID baseline. Easy examples therefore are not useless; they can provide stable optimization signal.")

    # 13 — annotation quality
    title = "Hard-to-learn regions are enriched for annotation problems"
    slide = deck.slide(title, section="Paper", citation="Swayamdipta et al. (2020), §5 and Figure 4.")
    steps = [
        ("1", "Flip labels for 1% of easy WinoGrande examples", PAPER),
        ("2", "Retrain and recompute training dynamics", TEAL),
        ("3", "Flipped examples move toward lower confidence", RED),
    ]
    for i, (num, text, color) in enumerate(steps):
        x = .72 + i * 4.18
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(1.55), Inches(.55), Inches(.55))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color); shape.line.color.rgb = rgb(color)
        textbox(slide, x, 1.62, .55, .25, num, 13, WHITE, bold=True, align=PP_ALIGN.CENTER, fill=color)
        textbox(slide, x + .7, 1.48, 3.15, .85, text, 15, TEXT, bold=True)
        if i < 2: textbox(slide, x + 3.66, 1.62, .42, .3, "→", 21, MUTED, bold=True, align=PP_ALIGN.CENTER)
    line(slide, .78, 2.72, 12.32, 2.72, GRID, 1)
    textbox(slide, .9, 3.08, 5.3, .34, "Human reannotation of predicted groups", 18, TEXT, bold=True)
    metric_table(slide, .9, 3.55, 5.3, 2.25,
                 ["Dataset", "Predicted noisy", "Predicted clean"],
                 [("WinoGrande", "67%", "13%"), ("SNLI", "76%", "4%")],
                 widths=[1.7, 1.75, 1.75], paper_cols=(1, 2), font_size=13)
    textbox(slide, 6.78, 3.08, 5.0, .34, "Responsible interpretation", 18, OURS, bold=True)
    bullet_list(slide, 6.78, 3.62, 5.05, 2.15, [
        "Low confidence is a useful review signal.",
        "The region also contains legitimate difficult cases.",
        "Hard-to-learn does not mean mislabeled.",
    ], size=15, bullet_color=OURS, gap=12)
    textbox(slide, 1.0, 6.28, 11.25, .38, "Data maps prioritize inspection; they do not replace annotation review.", 18, TEXT, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "The authors injected controlled noise by flipping labels for one percent of easy WinoGrande examples. After retraining, the corrupted examples shifted toward lower confidence. A simple confidence-based classifier then identified candidate issues. In human review, 67 percent versus 13 percent of WinoGrande’s predicted noisy and clean groups were mislabeled or ambiguous; on SNLI the contrast was 76 versus 4 percent. This shows enrichment, not perfect diagnosis.")

    # 14 — chosen result
    title = "The selected SNLI experiment tests data efficiency and generalization"
    slide = deck.slide(title, section="Paper", citation="Published values: Swayamdipta et al. (2020), Table 3; best of 3 random seeds.")
    textbox(slide, .78, 1.42, 5.65, .38, "IN-DISTRIBUTION (ID)", 11, PAPER, bold=True)
    textbox(slide, .78, 1.84, 5.65, .6, "SNLI test set", 25, TEXT, bold=True)
    textbox(slide, .78, 2.55, 5.65, .75, "Same collection process and task distribution as training.", 17, MUTED)
    line(slide, 6.58, 1.38, 6.58, 3.45, GRID, 1)
    textbox(slide, 6.98, 1.42, 5.4, .38, "OUT-OF-DISTRIBUTION (OOD)", 11, PAPER, bold=True)
    textbox(slide, 6.98, 1.84, 5.4, .6, "NLI Diagnostics", 25, TEXT, bold=True)
    textbox(slide, 6.98, 2.55, 5.15, .75, "1,104 hand-crafted examples probing lexical semantics, logic, predicate–argument structure, and knowledge.", 16, MUTED)
    metric_table(slide, 1.33, 3.75, 10.7, 2.28,
                 ["Training data", "Examples used", "SNLI test", "Diagnostics"],
                 [("Full", "100%", "92.0", "61.8"), ("Random", "33%", "91.3", "60.4"), ("Hard-to-learn", "33%", "91.8", "62.0"), ("Ambiguous", "33%", "92.2", "63.5")],
                 widths=[3.1, 2.3, 2.6, 2.6], paper_cols=(2, 3), font_size=13)
    textbox(slide, 1.15, 6.23, 11.05, .52, "Question: can one third of SNLI retain ID performance and improve diagnostic generalization?", 15.5, OURS, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    deck.add_notes(slide, title, "This is the concrete experiment we chose. ID means the ordinary SNLI test distribution. OOD means the NLI Diagnostics set, a small hand-crafted challenge set targeting multiple reasoning categories. The paper’s ambiguous third slightly exceeded the full-data ID score and improved Diagnostics by 1.7 points. The numerical ID difference is small; the interesting claim is data efficiency plus OOD behavior.")

    # 15 — protocol
    title = "Reproduction protocol follows the published SNLI selection"
    slide = deck.slide(title, section="Our reproduction")
    stages = [
        ("SNLI 1.0", "549,367 valid train pairs"),
        ("Paper coordinates", "confidence + variability"),
        ("Four conditions", "full · random · ambiguous · hard"),
        ("RoBERTa-large", "one seed per condition"),
        ("Two evaluations", "SNLI test · Diagnostics"),
    ]
    for i, (head, sub) in enumerate(stages):
        x = .52 + i * 2.57
        panel(slide, x, 1.62, 2.15, 1.42, fill=WHITE, border=OURS if i >= 3 else PAPER)
        textbox(slide, x + .15, 1.87, 1.85, .3, head, 15, TEXT, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, x + .15, 2.29, 1.85, .46, sub, 11, MUTED, align=PP_ALIGN.CENTER)
        if i < 4: textbox(slide, x + 2.16, 2.05, .38, .35, "→", 20, MUTED, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, .74, 3.57, 5.85, .34, "Selection definitions", 18, TEXT, bold=True)
    bullet_list(slide, .74, 4.04, 5.7, 1.92, [
        "Random: published 33% sampling procedure",
        "Ambiguous: highest variability",
        "Hard-to-learn: lowest confidence",
    ], size=15, bullet_color=PAPER, gap=11)
    textbox(slide, 6.88, 3.57, 5.7, .34, "Training settings", 18, TEXT, bold=True)
    bullet_list(slide, 6.88, 4.04, 5.55, 1.92, [
        "learning rate 1.07086 × 10⁻⁵",
        "effective batch 96 · maximum length 128",
        "up to 6 epochs · best validation checkpoint",
    ], size=15, bullet_color=OURS, gap=11)
    textbox(slide, 1.0, 6.42, 11.3, .28, "Published coordinates isolate the data-selection claim from map reconstruction differences.", 15, OURS, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "We retained every official SNLI example with a valid label and joined all 549,367 training IDs to the authors’ released coordinates. The random subset follows their published sampling seed and procedure. Ambiguous selects highest variability and hard-to-learn selects lowest confidence. Every model is RoBERTa-large with the recovered learning rate, effective batch 96, and best validation checkpoint.")

    # 16 — deviations
    title = "Hardware constraints changed implementation details"
    slide = deck.slide(title, section="Our reproduction", citation="Exact settings and evidence are documented in docs/original_experiment.md and results/run_metrics/.")
    rows = [
        ("GPU", "Quadro RTX 8000", "RTX 4070 Ti SUPER · 16 GB"),
        ("Batch", "physical batch 96", "32 × 3 gradient accumulation = 96"),
        ("Precision", "2020 training stack", "BF16 · modern PyTorch/Transformers"),
        ("Padding", "length 128", "dynamic padding · max length 128"),
        ("Reporting", "best of 3 seeds", "one available seed: 93078"),
        ("Epochs", "paper says 6; config early stopping", "best checkpoints; 4–6 epochs"),
    ]
    metric_table(slide, .72, 1.4, 11.9, 4.55, ["Aspect", "Original paper", "Our reproduction"], rows,
                 widths=[2.1, 4.6, 5.1], paper_cols=(1,), ours_cols=(2,), font_size=12.5)
    textbox(slide, .88, 6.15, 11.55, .62, "These facts make exact equality unlikely while leaving the subset comparison scientifically interpretable.", 15.5, TEXT, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    deck.add_notes(slide, title, "The original physical batch of 96 does not fit a 16 GB card, so we used 32 examples with three accumulation steps, preserving the effective batch. BF16 and dynamic padding reduce memory. The software stack is modern. Most importantly, the paper reports the best of three seeds, while the released configuration exposes one seed. Our comparisons therefore test whether the reported pattern appears, not whether every decimal is identical.")

    # 17 — ID
    title = "In-distribution accuracy closely matches the paper"
    slide = deck.slide(title, section="Our reproduction")
    add_image_fit(slide, GENERATED / "id_paired_dots.png", .58, 1.32, 8.0, 5.25)
    metric_table(slide, 8.6, 1.52, 3.95, 3.48,
                 ["Subset", "Paper", "Ours", "Δ"],
                 [(r.label.replace(" 33%", ""), f"{r.paper_snli:.1f}", f"{r.ours_snli:.2f}", f"{r.ours_snli-r.paper_snli:+.2f}") for _, r in metrics.iterrows()],
                 widths=[1.25, .85, .9, .85], paper_cols=(1,), ours_cols=(2,), font_size=11.5)
    panel(slide, 8.75, 5.35, 3.65, 1.05, fill="F1E8E2", border=OURS)
    textbox(slide, 8.98, 5.55, 3.18, .28, "Central reproduction result", 10, OURS, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 8.98, 5.9, 3.18, .28, "Ambiguous: 92.19 vs 92.2", 16, TEXT, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, .9, 6.62, 7.5, .22, "Axis is deliberately labeled as zoomed; direct values and differences are shown.", 10, MUTED)
    deck.add_notes(slide, title, "On SNLI test, full and random are about half a point above the paper. The ambiguous subset is essentially exact: 92.19 versus 92.2. It also beats our random third by 0.38 points and comes within 0.34 of our full-data run. Hard-to-learn is the largest ID disagreement, 1.38 points below the paper. The plotted axis is zoomed and explicitly labeled; the table keeps the size of each difference visible.")

    # 18 — OOD
    title = "The ambiguous subset also leads on NLI Diagnostics"
    slide = deck.slide(title, section="Our reproduction")
    add_image_fit(slide, GENERATED / "ood_paired_dots.png", .58, 1.32, 8.0, 5.25)
    metric_table(slide, 8.6, 1.52, 3.95, 3.48,
                 ["Subset", "Paper", "Ours", "Δ"],
                 [(r.label.replace(" 33%", ""), f"{r.paper_ood:.1f}", f"{r.ours_ood:.2f}", f"{r.ours_ood-r.paper_ood:+.2f}") for _, r in metrics.iterrows()],
                 widths=[1.25, .85, .9, .85], paper_cols=(1,), ours_cols=(2,), font_size=11.5)
    panel(slide, 8.75, 5.35, 3.65, 1.05, fill="F1E8E2", border=OURS)
    textbox(slide, 8.98, 5.55, 3.18, .28, "Our Diagnostics ranking", 10, OURS, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 8.98, 5.83, 3.18, .52, "Ambiguous  >  Full\n>  Hard  >  Random", 11.5, TEXT, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, .9, 6.62, 7.5, .22, "NLI Diagnostics contains 1,104 labeled examples; one seed can move several examples.", 10, MUTED)
    deck.add_notes(slide, title, "On Diagnostics, every reproduction score is within 1.65 points of the published value. Ambiguous remains best at 64.31, 1.54 points above our full run and 2.26 above our random third. The ranking supports the paper’s main OOD pattern. The random run is our largest positive difference. Diagnostics is only 1,104 examples, so a small number of predictions and seed differences can move the percentage.")

    # 19 — map comparison
    title = "Reproduced maps agree more on confidence than variability"
    slide = deck.slide(title, section="Our reproduction")
    add_image_fit(slide, RESULTS / "figures/snli_data_map.png", .52, 1.36, 5.0, 4.1)
    add_image_fit(slide, RESULTS / "figures/confidence_comparison.png", 5.45, 1.43, 3.25, 2.65)
    add_image_fit(slide, RESULTS / "figures/variability_comparison.png", 8.75, 1.43, 3.25, 2.65)
    stats = [
        ("Map coverage", f"{100*map_stats['coverage']:.2f}%", OURS),
        ("Confidence Pearson / Spearman", f"{map_stats['confidence']['pearson']:.3f} / {map_stats['confidence']['spearman']:.3f}", TEAL),
        ("Variability Pearson / Spearman", f"{map_stats['variability']['pearson']:.3f} / {map_stats['variability']['spearman']:.3f}", AMBER),
        ("Ambiguous-set overlap", f"{100*map_stats['ambiguous']['overlap']:.1f}%", PAPER),
    ]
    for i, (label, value, color) in enumerate(stats):
        x = .7 + i * 3.05
        textbox(slide, x, 4.72, 2.7, .25, label, 9.5, MUTED, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, x, 5.08, 2.7, .45, value, 22, color, bold=True, align=PP_ALIGN.CENTER)
    panel(slide, 1.2, 5.85, 10.9, .72, fill=WHITE, border=RED)
    textbox(slide, 1.45, 6.01, 10.4, .38, "Individual agreement: the paper’s “girl → boy” hard example is also among our lowest-confidence examples.", 15, TEXT, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, .76, 6.78, 11.8, .16, "Coverage excludes 498 examples lacking exactly one observation in every logged epoch; no values were imputed.", 8.5, MUTED, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "We also reconstructed the map from our own training dynamics. Coverage is 99.91 percent because 498 examples lacked exactly one observation in every logged epoch and were excluded rather than imputed. Confidence correlates strongly with the released map: Pearson 0.944. Variability is less stable but still substantial at 0.708. Seventy-five percent of ambiguous selections overlap. Qualitatively, the exact girl-versus-boy example highlighted as hard in the paper also appears among our lowest-confidence examples.")

    # 20 — synthesis/conclusion
    title = "The selection result reproduces; map details remain sensitive"
    slide = deck.slide(title, section="Synthesis")
    textbox(slide, .76, 1.43, 5.65, .36, "SUPPORTED BY OUR EVIDENCE", 11, GREEN, bold=True)
    bullet_list(slide, .76, 1.95, 5.7, 2.7, [
        "Ambiguous 33% matches paper ID accuracy: 92.19 vs 92.2.",
        "Ambiguous remains best on Diagnostics: 64.31.",
        "Confidence structure transfers strongly between maps.",
        "One third of the data approaches full-data performance.",
    ], size=15.5, bullet_color=GREEN, gap=12)
    textbox(slide, 6.85, 1.43, 5.7, .36, "LESS STABLE OR DIFFERENT", 11, RED, bold=True)
    bullet_list(slide, 6.85, 1.95, 5.55, 2.7, [
        "Hard-subset ID is 1.38 points below the paper.",
        "Variability agrees less strongly than confidence.",
        "The paper reports best of 3 seeds; we have one.",
        "Hardware, precision, and software stacks differ.",
    ], size=15.5, bullet_color=RED, gap=12)
    line(slide, .82, 4.9, 12.28, 4.9, GRID, 1)
    textbox(slide, 1.0, 5.12, 11.25, .68, "Training histories expose structure that final accuracy alone cannot show.", 20, TEXT, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(slide, 1.18, 5.88, 10.9, .78, "In this SNLI reproduction, selecting high-variability examples preserved accuracy and strengthened diagnostic generalization, while exact map coordinates remained sensitive to training details.", 15, OURS, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(slide, .82, 6.82, 2.0, .2, "Questions", 9, MUTED)
    textbox(slide, 4.25, 6.78, 8.1, .22, "github.com/JuliaMazine/dataset-cartography-reproduction", 8.5, MUTED, align=PP_ALIGN.RIGHT)
    deck.add_notes(slide, title, "The evidence supports the paper’s central SNLI selection result: the ambiguous third nearly matches full-data ID accuracy and leads on Diagnostics. The map’s confidence structure is robust, while variability and the hard-subset score are more sensitive. The broader lesson is that per-example training histories can diagnose usefulness, uncertainty, and possible annotation issues, but all region names remain conditional on the model and training process. That is the conclusion I would defend from this reproduction.")

    # Appendix A1 — exact settings
    title = "Exact reproduction hyperparameters"
    slide = deck.slide(title, section="Appendix", appendix=True)
    rows = [
        ("Model", "roberta-large"), ("Seed", "93078"), ("Learning rate", "1.0708609960508476 × 10⁻⁵"),
        ("Optimizer", "AdamW; ε=1e-8; weight decay 0; max grad norm 1"),
        ("Schedule", "linear decay; zero warmup"), ("Maximum length", "128 tokens"),
        ("Batch", "32 physical × 3 accumulation = 96 effective"), ("Precision", "BF16"),
        ("Checkpoint", "highest SNLI validation accuracy"), ("Epoch cap", "6; patience 3"),
    ]
    metric_table(slide, 1.15, 1.45, 11.0, 4.95, ["Setting", "Value"], rows, widths=[3.1, 7.9], ours_cols=(1,), font_size=12)
    deck.add_notes(slide, title, "Backup slide with the full settings that materially affect comparison. The recovered settings are documented in the repository and the raw run metrics.")

    # Appendix A2 — individual runs
    title = "Individual run metrics"
    slide = deck.slide(title, section="Appendix", appendix=True)
    rows = []
    for _, r in metrics.iterrows():
        raw = json.loads(next((RESULTS / "run_metrics").glob(f"{r['subset']}*.json")).read_text())
        rows.append((r.label, f"{int(r.train_examples):,}", f"{raw['epochs']:.2f}", f"{raw['runtime_minutes']:.1f} min", f"{raw['validation_accuracy']:.2f}", f"{raw['id_accuracy']:.2f}", f"{raw['ood_accuracy']:.2f}"))
    metric_table(slide, .55, 1.5, 12.25, 3.3,
                 ["Run", "Train N", "Epochs", "Runtime", "Validation", "SNLI test", "Diagnostics"], rows,
                 widths=[2.0, 1.55, 1.3, 1.5, 1.85, 1.85, 2.1], ours_cols=(4, 5, 6), font_size=11.5)
    textbox(slide, .82, 5.25, 11.7, .45, "All runs use one seed. Epoch counts differ because early stopping selected the best validation checkpoint.", 16, TEXT, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, .82, 6.0, 11.7, .4, "Raw JSON files: results/run_metrics/", 13, MUTED, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "Backup slide containing every individual measured run rather than rounded headline values. There is no hidden aggregation across seeds.")

    # Appendix A3 — subset construction
    title = "Subset construction"
    slide = deck.slide(title, section="Appendix", appendix=True)
    defs = [
        ("Random 33%", "181,291 examples", "Pandas sample with the recovered subset seed 725862.", PAPER),
        ("Ambiguous 33%", "182,335 examples", "Sort released coordinates by highest variability.", AMBER),
        ("Hard-to-learn 33%", "182,335 examples", "Sort released coordinates by lowest confidence.", RED),
        ("Full", "549,367 examples", "Every official SNLI training example with a valid label.", GREEN),
    ]
    for i, (name, count, desc, color) in enumerate(defs):
        y = 1.45 + i * 1.24
        pill(slide, .85, y, 1.85, name, color, fill=BG)
        textbox(slide, 3.05, y + .02, 2.0, .32, count, 15, color, bold=True)
        textbox(slide, 5.15, y - .04, 6.65, .52, desc, 16, TEXT)
    textbox(slide, 1.0, 6.52, 11.25, .3, "All coordinate joins are one-to-one and cover every valid training ID before selection.", 15, TEXT, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "Backup slide describing exact subset construction. The slight count difference between random and coordinate-based subsets follows the separate procedures in the original code.")

    # Appendix A4 — full table
    title = "Complete paper-versus-reproduction table"
    slide = deck.slide(title, section="Appendix", appendix=True)
    rows = [(r.label, f"{r.paper_snli:.1f}", f"{r.ours_snli:.2f}", f"{r.ours_snli-r.paper_snli:+.2f}",
             f"{r.paper_ood:.1f}", f"{r.ours_ood:.2f}", f"{r.ours_ood-r.paper_ood:+.2f}") for _, r in metrics.iterrows()]
    metric_table(slide, .56, 1.48, 12.2, 3.4,
                 ["Subset", "Paper ID", "Our ID", "Δ ID", "Paper OOD", "Our OOD", "Δ OOD"], rows,
                 widths=[2.2, 1.55, 1.55, 1.45, 1.75, 1.7, 1.7], paper_cols=(1, 4), ours_cols=(2, 5), font_size=11.5)
    textbox(slide, .85, 5.28, 11.6, .42, "Paper values are best-of-three; reproduction values are one run with seed 93078.", 16, TEXT, bold=True, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "Backup slide containing the full numerical comparison and signed differences in percentage points.")

    # Appendix A5 — environment and artifacts
    title = "Reproduction environment and artifact locations"
    slide = deck.slide(title, section="Appendix", appendix=True)
    textbox(slide, .78, 1.42, 5.7, .35, "Environment", 19, TEXT, bold=True)
    bullet_list(slide, .78, 1.92, 5.65, 2.5, [
        "NVIDIA RTX 4070 Ti SUPER · 16 GB",
        "Python 3.12 · PyTorch 2.3",
        "Transformers 4.44–4.49 range",
        "CUDA · BF16 · dynamic padding",
    ], size=15, bullet_color=OURS, gap=11)
    textbox(slide, 6.82, 1.42, 5.55, .35, "Repository artifacts", 19, TEXT, bold=True)
    bullet_list(slide, 6.82, 1.92, 5.45, 2.75, [
        "results/reproduction_metrics.csv",
        "results/run_metrics/*.json",
        "results/map_comparison.json",
        "presentation/generate_presentation.py",
        "presentation/speaker_notes.md",
    ], size=14.5, bullet_color=PAPER, gap=10)
    panel(slide, 1.1, 5.16, 11.1, .92, fill=WHITE, border=TEAL)
    textbox(slide, 1.35, 5.39, 10.6, .4, "The deck can be regenerated without model checkpoints from the committed compact results and assets.", 17, TEAL, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.05, 6.48, 11.2, .26, "Paper: Swayamdipta et al., Dataset Cartography, EMNLP 2020 · ACL Anthology 2020.emnlp-main.746", 10, MUTED, align=PP_ALIGN.CENTER)
    deck.add_notes(slide, title, "Backup slide for reproducibility questions and artifact navigation.")

    deck.save()


if __name__ == "__main__":
    build_deck()
